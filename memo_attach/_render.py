"""블록 → HTML / DataFrame / 내보내기 변환."""

from __future__ import annotations

import base64
import html as _html
import io
import re
from pathlib import Path
from typing import Any, Dict, List, Optional

TABLE_CSS = """
<style>
.ma-wrap { overflow-x: auto; margin: 2px 0 6px 0; }
.ma-table {
  border-collapse: collapse;
  font-size: 13px;
  font-family: "Malgun Gothic", "Apple SD Gothic Neo", "Noto Sans KR", sans-serif;
  background: #fff;
  color: #111;
}
.ma-table td, .ma-table th {
  border: 1px solid #d0d3d8;
  padding: 3px 7px;
  white-space: pre-wrap;
  overflow-wrap: break-word;
  vertical-align: middle;
}
.ma-table th { background: #f2f3f5; font-weight: 600; }
.ma-cap { font-size: 12px; color: #6b7280; margin: 0 0 4px 0; }
.ma-rich {
  background: #fff; color: #111; border: 1px solid #e3e5e9; border-radius: 8px;
  padding: 10px 12px; overflow-x: auto; font-size: 14px;
}
.ma-rich table { border-collapse: collapse; }
.ma-rich td, .ma-rich th { border: 1px solid #d0d3d8; padding: 3px 7px; }
</style>
"""

_SCRIPT_RE = re.compile(r"<\s*(script|iframe|object|embed|link|meta)\b.*?(</\s*\1\s*>|>)", re.I | re.S)
_ON_ATTR_RE = re.compile(r"\son\w+\s*=\s*(\"[^\"]*\"|'[^']*'|[^\s>]+)", re.I)


def sanitize_html(raw: str) -> str:
    """저장된 서식 텍스트를 다시 한 번 걸러낸다(방어적)."""
    out = _SCRIPT_RE.sub("", raw or "")
    out = _ON_ATTR_RE.sub("", out)
    out = re.sub(r"javascript\s*:", "", out, flags=re.I)
    return out


# --------------------------------------------------------------------------- #
# 표
# --------------------------------------------------------------------------- #
def _cell_style(cell: Dict[str, Any]) -> str:
    s: List[str] = []
    if cell.get("bg"):
        s.append(f"background-color:{cell['bg']}")
    if cell.get("fg"):
        s.append(f"color:{cell['fg']}")
    if cell.get("b"):
        s.append("font-weight:700")
    if cell.get("i"):
        s.append("font-style:italic")
    if cell.get("u"):
        s.append("text-decoration:underline")
    if cell.get("al"):
        s.append(f"text-align:{cell['al']}")
    if cell.get("va"):
        s.append(f"vertical-align:{cell['va']}")
    if cell.get("fs"):
        s.append(f"font-size:{cell['fs']}px")
    for key, css in (("bt", "border-top"), ("br", "border-right"), ("bb", "border-bottom"), ("bl", "border-left")):
        if cell.get(key):
            s.append(f"{css}:{cell[key]}")
    return ";".join(s)


def table_to_html(payload: Dict[str, Any], caption: str = "") -> str:
    grid = payload.get("grid") or []
    widths = payload.get("widths") or []
    parts: List[str] = ['<div class="ma-wrap">']
    if caption:
        parts.append(f'<div class="ma-cap">{_html.escape(caption)}</div>')
    parts.append('<table class="ma-table">')
    # 열 너비는 colgroup(고정 폭) 대신 셀의 min-width 로 준다.
    # 그래야 ① 엑셀에서 넓게 잡아둔 열은 그 폭을 유지하고
    #      ② 내용이 더 길면 표가 알아서 늘어나 한글이 어색하게 잘리지 않는다.
    width_done = set()
    for row in grid:
        parts.append("<tr>")
        for col, cell in enumerate(row):
            if cell is None:
                continue
            tag = "th" if cell.get("th") else "td"
            attrs = ""
            if cell.get("cs"):
                attrs += f' colspan="{int(cell["cs"])}"'
            if cell.get("rs"):
                attrs += f' rowspan="{int(cell["rs"])}"'
            style = _cell_style(cell)
            if not cell.get("cs") and col not in width_done and col < len(widths) and widths[col]:
                width_done.add(col)
                style = (style + ";" if style else "") + f"min-width:{int(widths[col])}px"
            if style:
                attrs += f' style="{style}"'
            text = _html.escape(str(cell.get("v", "")))
            parts.append(f"<{tag}{attrs}>{text}</{tag}>")
        parts.append("</tr>")
    parts.append("</table></div>")
    return "".join(parts)


def table_to_rows(payload: Dict[str, Any]) -> List[List[str]]:
    """병합 셀을 값 복제로 펴서 2차원 문자열 배열로."""
    grid = payload.get("grid") or []
    if not grid:
        return []
    ncols = max((len(r) for r in grid), default=0)
    out = [["" for _ in range(ncols)] for _ in grid]
    for r, row in enumerate(grid):
        for c, cell in enumerate(row):
            if cell is None:
                continue
            v = str(cell.get("v", ""))
            rs = int(cell.get("rs", 1) or 1)
            cs = int(cell.get("cs", 1) or 1)
            for i in range(rs):
                for j in range(cs):
                    if r + i < len(out) and c + j < ncols:
                        out[r + i][c + j] = v
    return out


def table_to_dataframe(payload: Dict[str, Any], header: bool = True):
    import pandas as pd

    rows = table_to_rows(payload)
    if not rows:
        return pd.DataFrame()
    if header and len(rows) > 1:
        cols, body = rows[0], rows[1:]
        seen: Dict[str, int] = {}
        uniq = []
        for i, c in enumerate(cols):
            name = c.strip() or f"열{i + 1}"
            if name in seen:
                seen[name] += 1
                name = f"{name}_{seen[name]}"
            else:
                seen[name] = 0
            uniq.append(name)
        return pd.DataFrame(body, columns=uniq)
    return pd.DataFrame(rows)


def table_to_csv_bytes(payload: Dict[str, Any]) -> bytes:
    import csv

    buf = io.StringIO()
    writer = csv.writer(buf)
    for row in table_to_rows(payload):
        writer.writerow(row)
    return buf.getvalue().encode("utf-8-sig")  # 엑셀에서 한글 안 깨지게 BOM


def table_to_xlsx_bytes(payload: Dict[str, Any]) -> Optional[bytes]:
    """서식(배경색/굵기/병합/열너비)까지 살려 xlsx로. openpyxl 없으면 None."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except Exception:
        return None

    grid = payload.get("grid") or []
    widths = payload.get("widths") or []
    wb = Workbook()
    ws = wb.active
    ws.title = "붙여넣은 표"

    for r, row in enumerate(grid, start=1):
        for c, cell in enumerate(row, start=1):
            if cell is None:
                continue
            target = ws.cell(row=r, column=c, value=cell.get("v", ""))
            font_kw: Dict[str, Any] = {}
            if cell.get("b"):
                font_kw["bold"] = True
            if cell.get("i"):
                font_kw["italic"] = True
            if cell.get("u"):
                font_kw["underline"] = "single"
            if cell.get("fg"):
                font_kw["color"] = str(cell["fg"]).lstrip("#").upper()
            if cell.get("fs"):
                font_kw["size"] = cell["fs"]
            if font_kw:
                target.font = Font(**font_kw)
            if cell.get("bg"):
                rgb = str(cell["bg"]).lstrip("#").upper()
                if len(rgb) == 6:
                    target.fill = PatternFill("solid", fgColor=rgb)
            if cell.get("al") or cell.get("va"):
                al = cell.get("al")
                al = al if al in ("left", "center", "right", "justify") else None
                target.alignment = Alignment(horizontal=al, vertical=cell.get("va") or None, wrap_text=True)
            rs = int(cell.get("rs", 1) or 1)
            cs = int(cell.get("cs", 1) or 1)
            if rs > 1 or cs > 1:
                ws.merge_cells(start_row=r, start_column=c, end_row=r + rs - 1, end_column=c + cs - 1)

    for i, w in enumerate(widths, start=1):
        if w:
            ws.column_dimensions[get_column_letter(i)].width = max(6, min(80, float(w) / 7.0))

    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


# --------------------------------------------------------------------------- #
# 블록 전체 → HTML (메모 통째로 내보내기)
# --------------------------------------------------------------------------- #
def _img_data_uri(path: str) -> str:
    p = Path(path)
    if not p.exists():
        return ""
    mime = "image/png"
    if p.suffix.lower() in (".jpg", ".jpeg"):
        mime = "image/jpeg"
    elif p.suffix.lower() == ".gif":
        mime = "image/gif"
    elif p.suffix.lower() == ".webp":
        mime = "image/webp"
    return f"data:{mime};base64,{base64.b64encode(p.read_bytes()).decode('ascii')}"


def block_to_html(block: Dict[str, Any], embed_images: bool = True) -> str:
    kind = block["kind"]
    cap = block.get("caption") or ""
    if kind == "table":
        return table_to_html(block["payload"], cap)
    if kind == "image":
        src = _img_data_uri(block["file_path"]) if embed_images else Path(block["file_path"]).name
        capline = f'<div class="ma-cap">{_html.escape(cap)}</div>' if cap else ""
        return f'<div class="ma-wrap">{capline}<img src="{src}" style="max-width:100%;border:1px solid #e3e5e9;border-radius:6px"></div>'
    if kind == "html":
        capline = f'<div class="ma-cap">{_html.escape(cap)}</div>' if cap else ""
        return f'<div class="ma-wrap">{capline}<div class="ma-rich">{sanitize_html(block["payload"].get("html", ""))}</div></div>'
    text = _html.escape(block["payload"].get("text", ""))
    return f'<div class="ma-wrap"><pre style="white-space:pre-wrap;font-family:inherit;margin:0">{text}</pre></div>'


def memo_to_html(title: str, body_text: str, blocks: List[Dict[str, Any]]) -> str:
    """메모 본문 + 첨부 전체를 하나의 HTML 문서로."""
    parts = [
        "<!doctype html><html lang='ko'><head><meta charset='utf-8'>",
        f"<title>{_html.escape(title)}</title>",
        TABLE_CSS,
        "<style>body{font-family:'Malgun Gothic','Apple SD Gothic Neo','Noto Sans KR',sans-serif;"
        "max-width:1100px;margin:32px auto;padding:0 20px;color:#111;background:#fff;line-height:1.7}"
        "h1{border-bottom:2px solid #ff4b4b;padding-bottom:8px}"
        ".memo-body{white-space:pre-wrap;margin-bottom:24px}</style></head><body>",
        f"<h1>{_html.escape(title)}</h1>",
        f"<div class='memo-body'>{_html.escape(body_text or '')}</div>",
    ]
    for b in blocks:
        parts.append(block_to_html(b, embed_images=True))
    parts.append("</body></html>")
    return "".join(parts)


def memo_to_markdown(title: str, body_text: str, blocks: List[Dict[str, Any]]) -> str:
    lines = [f"# {title}", "", body_text or "", ""]
    for b in blocks:
        cap = b.get("caption") or ""
        if b["kind"] == "table":
            rows = table_to_rows(b["payload"])
            if cap:
                lines.append(f"**{cap}**")
            if rows:
                lines.append("| " + " | ".join(x.replace("|", "\\|") for x in rows[0]) + " |")
                lines.append("|" + "---|" * len(rows[0]))
                for r in rows[1:]:
                    lines.append("| " + " | ".join(x.replace("|", "\\|").replace("\n", "<br>") for x in r) + " |")
            lines.append("")
        elif b["kind"] == "image":
            lines.append(f"![{cap or '이미지'}]({Path(b['file_path']).name})")
            lines.append("")
        elif b["kind"] == "html":
            lines.append(b["payload"].get("plain") or "")
            lines.append("")
        else:
            lines.append(b["payload"].get("text", ""))
            lines.append("")
    return "\n".join(lines)
